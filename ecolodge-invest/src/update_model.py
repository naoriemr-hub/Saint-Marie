"""Injecte le prix live de la source 'feeds_model' dans le modèle Excel.

Écrit le prix dans Scenario Rachat!C5. Les formules se recalculent à l'ouverture
d'Excel (ou via LibreOffice : voir README).
"""
import json, sys
from pathlib import Path
import openpyxl, yaml

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "model/Ecolodge_Sainte_Marie_Business_Plan.xlsx"


def main() -> int:
    cfg = yaml.safe_load((ROOT / "config/sources.yaml").read_text())
    state = json.loads((ROOT / "data/listings.json").read_text())
    feed = next((s for s in cfg["sources"] if s.get("feeds_model")), None)
    if not feed:
        print("aucune source feeds_model"); return 0
    rec = state.get(feed["id"], {})
    price, status = rec.get("price"), rec.get("status")
    if not price or status != "actif":
        print(f"prix non mis à jour (price={price}, status={status})"); return 0
    wb = openpyxl.load_workbook(XLSX)
    ws = wb["Scenario Rachat"]
    old = ws["C5"].value
    if old == price:
        print("prix inchangé"); return 0
    ws["C5"] = price
    wb.save(XLSX)
    print(f"Scenario Rachat!C5 : {old} → {price}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
